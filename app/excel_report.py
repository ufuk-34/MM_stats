"""Grafikli Excel raporu.

Sayfalar:
  Rapor     - Dashboard görünümü: özet kutuları, öne çıkanlar ve Excel grafikleri (tek A4 sayfasına yazdırılır)
  Tablolar  - Grafiklerin dayandığı özet tablolar
  Kayıtlar  - Filtrelenmiş kayıtların tam listesi

Rapor ve Tablolar sayfalarındaki tüm sayılar "Kayıtlar" sayfasından formülle hesaplanır;
Excel dosya açıldığında hesaplar. Böylece rapor ile kayıt listesi her zaman birbirini tutar.
"""
from collections import Counter
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.legend import Legend
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.text import CharacterProperties, Font as DrawingFont, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

from .constants import (
    SOURCE_SYSTEM, SOURCE_USER, SOURCES, STATUS_CLOSED, STATUS_IN_REVIEW, STATUS_OPEN, STATUS_RESOLVED, STATUSES,
)
from .extensions import db
from .queries import by_category, by_project, by_user, ticket_query

FONT = "Arial"
NAVY = "1B2B44"
BLUE = "2A78D6"        # Sistem Kaynaklı / tek seri
ORANGE = "EB6834"      # Kullanıcı İşlemi Kaynaklı
GREEN = "1D7A4F"       # Çözülen
AMBER = "EDA100"       # Bekleyen
MUTED = "6B7686"
TILE_FILL = PatternFill("solid", fgColor="F3F5F8")
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
THIN = Side(style="thin", color="DDE2E9")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

RECORD_COLUMNS = [
    ("Kayıt No", 12), ("Tarih", 11), ("Saat", 7), ("Proje", 28), ("Talep Sahibi", 22),
    ("Telefon", 15), ("İletişim Kanalı", 14), ("Sorun Kaynağı", 24), ("Sorun Kategorisi", 20),
    ("Öncelik", 9), ("Açıklama", 50), ("Personel", 20), ("İşlem Süresi (dk)", 10),
    ("Sonuç / Yapılan İşlem", 50), ("Durum", 12),
]
# Kayıtlar sayfasındaki sütun harfleri
COL = {"no": "A", "date": "B", "project": "D", "source": "H", "category": "I",
       "staff": "L", "duration": "M", "status": "O"}

MAX_DAILY_POINTS = 31      # daha uzun aralıklar aylık gösterilir
MAX_CATEGORIES = 10


def _font(size=10, bold=False, color="000000"):
    return Font(name=FONT, size=size, bold=bold, color=color)


def _set(ws, ref, value, **style):
    """Hücreye değer yazar. Metin değerler asla formül olarak yorumlanmaz (formül enjeksiyonu)."""
    cell = ws[ref] if isinstance(ref, str) else ws.cell(row=ref[0], column=ref[1])
    cell.value = value
    if isinstance(value, str) and not style.pop("formula", False):
        cell.data_type = "s"
    for key, val in style.items():
        setattr(cell, key, val)
    return cell


def _f(ws, ref, formula, **style):
    return _set(ws, ref, formula, formula=True, **style)


def _header_row(ws, row, titles, start_col=1):
    for i, title in enumerate(titles):
        c = ws.cell(row=row, column=start_col + i, value=title)
        c.fill, c.font = HEADER_FILL, _font(bold=True, color="FFFFFF")
        c.alignment = Alignment(vertical="center", wrap_text=True)


class _Ranges:
    """Kayıtlar sayfasındaki sütun aralıklarını formül metni olarak verir."""

    def __init__(self, row_count):
        self.last = max(row_count + 1, 2)

    def __call__(self, key):
        col = COL[key]
        return f"'Kayıtlar'!${col}$2:${col}${self.last}"


def _staff_labels(tickets, staff_rows):
    """Aynı ada sahip iki personel varsa ayırt etmek için kullanıcı adı eklenir."""
    users = {t.created_by.id: t.created_by for t in tickets}
    names = Counter(u.full_name for u in users.values())
    labels = {uid: (u.full_name if names[u.full_name] == 1 else f"{u.full_name} ({u.username})")
              for uid, u in users.items()}
    for row in staff_rows:
        labels.setdefault(row["id"], row["name"])
    return labels


def _time_buckets(filters, tickets):
    """Zaman grafiği için günler (kısa aralık) veya aylar (uzun aralık)."""
    dates = [t.created_at.date() for t in tickets]
    start = filters.date_from or (min(dates) if dates else date.today())
    end = filters.date_to or (max(dates) if dates else date.today())
    end = min(end, max(date.today(), start))
    if (end - start).days < MAX_DAILY_POINTS:
        return "daily", [start + timedelta(days=i) for i in range((end - start).days + 1)]
    months, d = [], start.replace(day=1)
    while d <= end:
        months.append(d)
        d = (d + timedelta(days=32)).replace(day=1)
    return "monthly", months


# --------------------------------------------------------------------------- #
# Sayfalar
# --------------------------------------------------------------------------- #

def _records_sheet(ws, tickets, staff_label):
    ws.title = "Kayıtlar"
    _header_row(ws, 1, [c[0] for c in RECORD_COLUMNS])
    for r, t in enumerate(tickets, start=2):
        values = [
            t.ticket_no, t.created_at.date(), t.created_at.strftime("%H:%M"), t.project.name,
            t.requester_name, t.requester_phone, t.channel_label, t.source_label, t.category.name,
            t.priority_label, t.description, staff_label[t.created_by_id], t.duration_minutes,
            t.resolution, t.status_label,
        ]
        for c, value in enumerate(values, start=1):
            cell = _set(ws, (r, c), value)
            cell.font = _font()
        ws.cell(row=r, column=2).number_format = "dd.mm.yyyy"
        ws.cell(row=r, column=11).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=14).alignment = Alignment(wrap_text=True, vertical="top")
    for i, (_, width) in enumerate(RECORD_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(RECORD_COLUMNS))}{max(len(tickets) + 1, 2)}"
    ws.page_setup.orientation = "landscape"
    ws.print_title_rows = "1:1"


def _tables_sheet(ws, R, projects, categories, staff_rows, staff_label, buckets_kind, buckets):
    """Özet tablolar. Her tablonun konumunu (grafikler ve Rapor sayfası için) döndürür."""
    ws.column_dimensions["A"].width = 34
    for col in "BCDEF":
        ws.column_dimensions[col].width = 15
    pos = {}

    def title(row, text):
        _set(ws, (row, 1), text, font=_font(12, True, NAVY))

    # Genel sayılar
    title(1, "Genel Özet")
    sys_label, usr_label = SOURCES[SOURCE_SYSTEM], SOURCES[SOURCE_USER]
    summary_rows = [
        ("total", "Toplam destek kaydı", f"=COUNTA({R('no')})", "0"),
        ("system", sys_label, f'=COUNTIF({R("source")},"{sys_label}")', "0"),
        ("user", usr_label, f'=COUNTIF({R("source")},"{usr_label}")', "0"),
        ("done", "Çözülen (Çözüldü + Kapatıldı)",
         f'=COUNTIF({R("status")},"{STATUSES[STATUS_RESOLVED]}")+COUNTIF({R("status")},"{STATUSES[STATUS_CLOSED]}")', "0"),
        ("pending", "Açık / Bekleyen",
         f'=COUNTIF({R("status")},"{STATUSES[STATUS_OPEN]}")+COUNTIF({R("status")},"{STATUSES[STATUS_IN_REVIEW]}")', "0"),
        ("avg", "Ortalama işlem süresi (dk)", f'=IFERROR(ROUND(AVERAGE({R("duration")}),1),0)', "0.0"),
        ("sum", "Toplam işlem süresi (dk)", f'=SUM({R("duration")})', "0"),
    ]
    for i, (key, label, formula, fmt) in enumerate(summary_rows, start=2):
        _set(ws, (i, 1), label, font=_font())
        _f(ws, (i, 2), formula, font=_font(bold=True), number_format=fmt)
        pos[key] = f"'Tablolar'!$B${i}"
    total = pos["total"]
    for key, row in (("system_pct", 3), ("user_pct", 4)):
        _f(ws, (row, 3), f"=IF({total}=0,0,B{row}/{total})", font=_font(color=MUTED), number_format="0.0%")
        pos[key] = f"'Tablolar'!$C${row}"

    def name_table(start, heading, columns, names, key_range, extra=None):
        """Ada göre sayım tablosu: Ad | Kayıt | % (+ ek sütunlar)."""
        title(start, heading)
        _header_row(ws, start + 1, columns)
        first = start + 2
        for i, name in enumerate(names):
            r = first + i
            _set(ws, (r, 1), name, font=_font())
            _f(ws, (r, 2), f"=SUMPRODUCT(--({key_range}=$A{r}))", font=_font())
            _f(ws, (r, 3), f"=IF({total}=0,0,B{r}/{total})", font=_font(color=MUTED), number_format="0.0%")
            if extra:
                extra(r)
        last = first + max(len(names), 1) - 1
        if not names:
            _set(ws, (first, 1), "Kayıt yok", font=_font(color=MUTED))
            ws.cell(row=first, column=2, value=0)
        return first, last

    row = len(summary_rows) + 4
    pos["projects"] = name_table(row, "Projelere Göre", ["Proje", "Kayıt", "Oran"],
                                 [p["name"] for p in projects], R("project"))

    row = pos["projects"][1] + 3
    title(row, "Sorun Kaynağına Göre")
    _header_row(ws, row + 1, ["Sorun Kaynağı", "Kayıt", "Oran"])
    for i, (label, key) in enumerate(((sys_label, "system"), (usr_label, "user"))):
        r = row + 2 + i
        _set(ws, (r, 1), label, font=_font())
        _f(ws, (r, 2), f"={pos[key]}", font=_font())
        _f(ws, (r, 3), f"=IF({total}=0,0,B{r}/{total})", font=_font(color=MUTED), number_format="0.0%")
    pos["source"] = (row + 2, row + 3)

    row = pos["source"][1] + 3
    pos["categories"] = name_table(row, "Sorun Kategorilerine Göre (en sık)", ["Kategori", "Kayıt", "Oran"],
                                   [c["name"] for c in categories[:MAX_CATEGORIES]], R("category"))

    # Zaman tablosu (gün veya ay)
    row = pos["categories"][1] + 3
    title(row, "Günlere Göre" if buckets_kind == "daily" else "Aylara Göre")
    _header_row(ws, row + 1, ["Tarih" if buckets_kind == "daily" else "Ay", sys_label, usr_label, "Toplam"])
    first = row + 2
    for i, d in enumerate(buckets):
        r = first + i
        cell = ws.cell(row=r, column=1, value=d)
        cell.font = _font()
        if buckets_kind == "daily":
            cell.number_format = "dd.mm"
            cond = f"{R('date')},$A{r}"
        else:
            cell.number_format = "mm.yyyy"
            cond = f'{R("date")},">="&$A{r},{R("date")},"<"&DATE(YEAR($A{r}),MONTH($A{r})+1,1)'
        _f(ws, (r, 2), f'=COUNTIFS({cond},{R("source")},"{sys_label}")', font=_font())
        _f(ws, (r, 3), f'=COUNTIFS({cond},{R("source")},"{usr_label}")', font=_font())
        _f(ws, (r, 4), f"=B{r}+C{r}", font=_font(bold=True))
    pos["time"] = (first, first + len(buckets) - 1)

    # Personel iş yükü
    row = pos["time"][1] + 3
    title(row, "Personel İş Yükü")
    _header_row(ws, row + 1, ["Personel", "Kayıt", "Çözülen", "Bekleyen", "Toplam Süre (dk)"])
    first = row + 2
    staff, status = R("staff"), R("status")
    done_a, done_b = STATUSES[STATUS_RESOLVED], STATUSES[STATUS_CLOSED]
    pend_a, pend_b = STATUSES[STATUS_OPEN], STATUSES[STATUS_IN_REVIEW]
    for i, u in enumerate(staff_rows):
        r = first + i
        _set(ws, (r, 1), staff_label[u["id"]], font=_font())
        _f(ws, (r, 2), f"=SUMPRODUCT(--({staff}=$A{r}))", font=_font())
        _f(ws, (r, 3), f'=SUMPRODUCT(({staff}=$A{r})*(({status}="{done_a}")+({status}="{done_b}")))', font=_font())
        _f(ws, (r, 4), f'=SUMPRODUCT(({staff}=$A{r})*(({status}="{pend_a}")+({status}="{pend_b}")))', font=_font())
        _f(ws, (r, 5), f"=SUMPRODUCT(({staff}=$A{r})*{R('duration')})", font=_font())
    if not staff_rows:
        _set(ws, (first, 1), "Kayıt yok", font=_font(color=MUTED))
    pos["staff"] = (first, first + max(len(staff_rows), 1) - 1)
    return pos


def _tile(ws, col, label, value_formula, sub_formula=None, value_format="0", sub_format=None):
    """Rapor sayfasında 2 sütun genişliğinde özet kutusu (satır 6-8)."""
    a, b = get_column_letter(col), get_column_letter(col + 1)
    for r in (6, 7, 8):
        ws.merge_cells(f"{a}{r}:{b}{r}")
        for c in (col, col + 1):
            ws.cell(row=r, column=c).fill = TILE_FILL
            ws.cell(row=r, column=c).border = Border(
                left=THIN if c == col else None, right=THIN if c == col + 1 else None,
                top=THIN if r == 6 else None, bottom=THIN if r == 8 else None)
    _set(ws, f"{a}6", label, font=_font(9, color=MUTED), alignment=Alignment(horizontal="center"))
    _f(ws, f"{a}7", value_formula, font=_font(20, True, NAVY), number_format=value_format,
       alignment=Alignment(horizontal="center", vertical="center"))
    if sub_formula:
        _f(ws, f"{a}8", sub_formula, font=_font(9, color=MUTED), number_format=sub_format or "General",
           alignment=Alignment(horizontal="center"))


def _chars(size, bold=False, color="404040"):
    """Grafik yazı tipi (boyut: punto x 100)."""
    return CharacterProperties(sz=size, b=bold, solidFill=color, latin=DrawingFont(typeface=FONT))


def _rich(size, bold=False, color="404040"):
    cp = _chars(size, bold, color)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _title(text):
    cp = _chars(1100, True, NAVY)
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])
    return Title(tx=Text(rich=RichText(p=[para])), overlay=False)


def _style_bar(chart, title, color=None, horizontal=False):
    chart.title = _title(title)
    chart.style = 10
    chart.legend = None
    chart.x_axis.txPr = _rich(800)
    chart.y_axis.txPr = _rich(800)
    chart.gapWidth = 60
    chart.x_axis.delete = False      # openpyxl 3.1: eksenler Excel'de görünsün
    chart.y_axis.delete = False
    chart.y_axis.majorGridlines = None
    chart.y_axis.numFmt = "0"
    if horizontal:
        chart.type = "bar"
        chart.x_axis.scaling.orientation = "maxMin"   # ilk kalem üstte
    else:
        chart.type = "col"
    if color:
        for s in chart.series:
            s.graphicalProperties.solidFill = color
            s.graphicalProperties.line.solidFill = color


def _value_labels(series):
    series.dLbls = DataLabelList()
    series.dLbls.showVal = True
    series.dLbls.txPr = _rich(800, True)
    for attr in ("showSerName", "showCatName", "showLegendKey", "showPercent"):
        setattr(series.dLbls, attr, False)


def _report_sheet(ws, R, pos, org_name, period_text, filter_text, buckets_kind):
    ws.sheet_view.showGridLines = False
    for c in range(1, 13):
        ws.column_dimensions[get_column_letter(c)].width = 11.5

    ws.merge_cells("A1:L1")
    _set(ws, "A1", "TEKNİK DESTEK RAPORU", font=_font(18, True, "FFFFFF"),
         fill=PatternFill("solid", fgColor=NAVY), alignment=Alignment(horizontal="left", vertical="center", indent=1))
    ws.row_dimensions[1].height = 34
    ws.merge_cells("A2:L2")
    _set(ws, "A2", f"{org_name}  ·  {period_text}", font=_font(11, True, NAVY), alignment=Alignment(indent=1))
    ws.merge_cells("A3:L3")
    _set(ws, "A3", f"{filter_text}  ·  Oluşturma: {datetime.now():%d.%m.%Y %H:%M}",
         font=_font(9, color=MUTED), alignment=Alignment(indent=1))

    t = pos
    _tile(ws, 1, "Toplam Kayıt", f"={t['total']}")
    _tile(ws, 3, "Çözülen", f"={t['done']}",
          f"=IF({t['total']}=0,\"\",{t['done']}/{t['total']})", sub_format='0.0%" çözüm oranı"')
    _tile(ws, 5, "Açık / Bekleyen", f"={t['pending']}")
    _tile(ws, 7, "Sistem Kaynaklı", f"={t['system']}", f"={t['system_pct']}", sub_format="0.0%")
    _tile(ws, 9, "Kullanıcı İşlemi Kaynaklı", f"={t['user']}", f"={t['user_pct']}", sub_format="0.0%")
    _tile(ws, 11, "Ort. İşlem Süresi", f"={t['avg']}", f"={t['sum']}",
          value_format='0.0" dk"', sub_format='"Toplam "0" dk"')
    for r, h in ((6, 16), (7, 30), (8, 16)):
        ws.row_dimensions[r].height = h

    # Öne çıkanlar: rapordan tek bakışta okunacak cümleler
    p1, p2 = t["projects"]
    c1, c2 = t["categories"]
    pr = f"'Tablolar'!$A${p1}:$A${p2}", f"'Tablolar'!$B${p1}:$B${p2}"
    ca = f"'Tablolar'!$A${c1}:$A${c2}", f"'Tablolar'!$B${c1}:$B${c2}"
    highlight = (
        f'=IF({t["total"]}=0,"Bu dönemde kayıt bulunmuyor.",'
        f'"En çok destek talebi: "&INDEX({pr[0]},MATCH(MAX({pr[1]}),{pr[1]},0))&" ("&MAX({pr[1]})&" kayıt)'
        f'   ·   En sık sorun: "&INDEX({ca[0]},MATCH(MAX({ca[1]}),{ca[1]},0))&" ("&MAX({ca[1]})&" kayıt)'
        f'   ·   Sistem kaynaklı oranı: "&TEXT({t["system_pct"]},"0%"))'
    )
    ws.merge_cells("A10:L10")
    _f(ws, "A10", highlight, font=_font(10, True, NAVY), fill=PatternFill("solid", fgColor="E8EFF8"),
       alignment=Alignment(indent=1, vertical="center"))
    ws.row_dimensions[10].height = 22

    tables = ws.parent["Tablolar"]

    # Grafik 1: Projelere göre
    ch = BarChart()
    ch.add_data(Reference(tables, min_col=2, min_row=p1 - 1, max_row=p2), titles_from_data=True)
    ch.set_categories(Reference(tables, min_col=1, min_row=p1, max_row=p2))
    _style_bar(ch, "Projelere Göre Kayıtlar", BLUE, horizontal=True)
    _value_labels(ch.series[0])
    ch.width, ch.height = 13.4, 7.2
    ws.add_chart(ch, "A12")

    # Grafik 2: Sorun kaynağı (halka)
    s1, s2 = t["source"]
    dn = DoughnutChart(holeSize=55)
    dn.add_data(Reference(tables, min_col=2, min_row=s1 - 1, max_row=s2), titles_from_data=True)
    dn.set_categories(Reference(tables, min_col=1, min_row=s1, max_row=s2))
    dn.title = _title("Sorun Kaynağı Dağılımı")
    dn.style = 10
    for idx, color in enumerate((BLUE, ORANGE)):
        pt = DataPoint(idx=idx)
        pt.graphicalProperties.solidFill = color
        pt.graphicalProperties.line.solidFill = "FFFFFF"
        dn.series[0].dPt.append(pt)
    dn.series[0].dLbls = DataLabelList()
    dn.series[0].dLbls.showPercent = True
    for attr in ("showVal", "showSerName", "showCatName", "showLegendKey"):
        setattr(dn.series[0].dLbls, attr, False)
    dn.series[0].dLbls.txPr = _rich(900, True, "FFFFFF")
    dn.legend.position = "b"
    dn.legend.txPr = _rich(800)
    dn.width, dn.height = 13.4, 7.2
    ws.add_chart(dn, "G12")

    # Grafik 3: Günlere / aylara göre (kaynağa göre yığılmış)
    f1, f2 = t["time"]
    tc = BarChart()
    tc.add_data(Reference(tables, min_col=2, max_col=3, min_row=f1 - 1, max_row=f2), titles_from_data=True)
    tc.set_categories(Reference(tables, min_col=1, min_row=f1, max_row=f2))
    _style_bar(tc, "Günlere Göre Kayıtlar" if buckets_kind == "daily" else "Aylara Göre Kayıtlar")
    tc.grouping, tc.overlap = "stacked", 100
    tc.series[0].graphicalProperties.solidFill = BLUE
    tc.series[1].graphicalProperties.solidFill = ORANGE
    tc.legend = Legend(legendPos="t")
    tc.legend.txPr = _rich(800)
    tc.x_axis.number_format = "dd.mm" if buckets_kind == "daily" else "mm.yyyy"
    tc.width, tc.height = 26.9, 7.2
    ws.add_chart(tc, "A27")

    # Grafik 4: Sorun kategorileri
    k = BarChart()
    k.add_data(Reference(tables, min_col=2, min_row=c1 - 1, max_row=c2), titles_from_data=True)
    k.set_categories(Reference(tables, min_col=1, min_row=c1, max_row=c2))
    _style_bar(k, "En Sık Sorun Kategorileri", BLUE, horizontal=True)
    _value_labels(k.series[0])
    k.width, k.height = 13.4, 9.0
    ws.add_chart(k, "A42")

    # Grafik 5: Personel iş yükü (çözülen + bekleyen)
    u1, u2 = t["staff"]
    pc = BarChart()
    pc.add_data(Reference(tables, min_col=3, max_col=4, min_row=u1 - 1, max_row=u2), titles_from_data=True)
    pc.set_categories(Reference(tables, min_col=1, min_row=u1, max_row=u2))
    _style_bar(pc, "Personel İş Yükü", horizontal=True)
    pc.grouping, pc.overlap = "stacked", 100
    pc.series[0].graphicalProperties.solidFill = GREEN
    pc.series[1].graphicalProperties.solidFill = AMBER
    pc.legend = Legend(legendPos="b")
    pc.legend.txPr = _rich(800)
    pc.width, pc.height = 13.4, 9.0
    ws.add_chart(pc, "G42")

    # Yazdırma: tek A4 sayfası (dikey)
    ws.print_area = "A1:L60"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.oddFooter.center.text = "Sayfa &P / &N"


def build_report_workbook(filters, user, conds, stat_conds, org_name, period_text, filter_text):
    """Grafikli Excel raporunu oluşturur. (Workbook, kayıt sayısı) döndürür."""
    tickets = db.session.execute(ticket_query(conds)).unique().scalars().all()
    projects = by_project(stat_conds)
    categories = by_category(stat_conds)
    staff_rows = by_user(stat_conds)
    staff_label = _staff_labels(tickets, staff_rows)
    buckets_kind, buckets = _time_buckets(filters, tickets)

    wb = Workbook()
    wb._named_styles["Normal"].font = _font()
    wb.calculation = CalcProperties(fullCalcOnLoad=True)   # Excel açılışta tüm formülleri hesaplar

    report = wb.active
    report.title = "Rapor"
    tables = wb.create_sheet("Tablolar")
    records = wb.create_sheet("Kayıtlar")

    R = _Ranges(len(tickets))
    _records_sheet(records, tickets, staff_label)
    pos = _tables_sheet(tables, R, projects, categories, staff_rows, staff_label, buckets_kind, buckets)
    _report_sheet(report, R, pos, org_name, period_text, filter_text, buckets_kind)
    wb.active = 0
    return wb, len(tickets)
