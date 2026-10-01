"""Raporlar: filtrelenmiş kayıtların Excel'e aktarılması (yalnızca yönetici)."""
from datetime import date, datetime
from io import BytesIO

from flask import Blueprint, render_template, request, send_file
from flask_login import current_user
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .auth import admin_required
from .extensions import db
from .models import Category, Project, User
from .queries import (
    STATUS_GROUPS, TicketFilters, by_category, by_project, by_user, stat_conditions, summary,
    ticket_query,
)
from .constants import SOURCES, STATUSES

bp = Blueprint("reports", __name__, url_prefix="/reports")

EXPORT_COLUMNS = [
    ("Kayıt No", 12), ("Tarih", 11), ("Saat", 7), ("Proje", 28), ("Talep Sahibi", 22),
    ("Telefon", 15), ("İletişim Kanalı", 14), ("Sorun Kaynağı", 24), ("Sorun Kategorisi", 20),
    ("Öncelik", 9), ("Açıklama", 50), ("Personel", 20), ("İşlem Süresi (dk)", 10),
    ("Sonuç / Yapılan İşlem", 50), ("Durum", 12),
]
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")


@bp.route("/")
@admin_required
def index():
    filters = TicketFilters.from_args(request.args)
    if not request.args:
        # Varsayılan: bu ayın kayıtları
        filters.date_from, filters.date_to = date.today().replace(day=1), date.today()
    count = summary(filters.conditions(current_user))["total"]
    return render_template(
        "reports/index.html", filters=filters, count=count, status_groups=STATUS_GROUPS,
        users=db.session.execute(db.select(User).order_by(User.full_name)).scalars().all(),
        projects=db.session.execute(db.select(Project).order_by(Project.sort_order, Project.id)).scalars().all(),
        categories=db.session.execute(db.select(Category).order_by(Category.sort_order, Category.name)).scalars().all(),
    )


def _write_row(ws, row_idx, values):
    for col_idx, value in enumerate(values, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        if isinstance(value, str):
            # Formül enjeksiyonuna karşı: "=" ile başlayan metin formül olarak yorumlanmasın
            cell.data_type = "s"


def _header(ws, row_idx, titles):
    for col_idx, title in enumerate(titles, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=title)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def filter_description(filters):
    parts = []
    if filters.date_from or filters.date_to:
        d1 = filters.date_from.strftime("%d.%m.%Y") if filters.date_from else "…"
        d2 = filters.date_to.strftime("%d.%m.%Y") if filters.date_to else "…"
        parts.append(f"Tarih: {d1} - {d2}")
    if filters.project_id:
        p = db.session.get(Project, filters.project_id)
        parts.append(f"Proje: {p.name if p else '-'}")
    if filters.user_id:
        u = db.session.get(User, filters.user_id)
        parts.append(f"Personel: {u.full_name if u else '-'}")
    if filters.source:
        parts.append(f"Sorun kaynağı: {SOURCES[filters.source]}")
    if filters.category_id:
        c = db.session.get(Category, filters.category_id)
        parts.append(f"Kategori: {c.name if c else '-'}")
    if filters.status:
        label = STATUS_GROUPS[filters.status][0] if filters.status in STATUS_GROUPS else STATUSES[filters.status]
        parts.append(f"Durum: {label}")
    if filters.q:
        parts.append(f"Arama: {filters.q}")
    if filters.archived:
        parts.append("Yalnızca arşivlenmiş kayıtlar")
    return "; ".join(parts) or "Filtre yok (tüm kayıtlar)"


def build_workbook(filters, user):
    conds = filters.conditions(user)
    tickets = db.session.execute(ticket_query(conds)).unique().scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Kayıtlar"
    _header(ws, 1, [c[0] for c in EXPORT_COLUMNS])
    for i, t in enumerate(tickets, start=2):
        _write_row(ws, i, [
            t.ticket_no, t.created_at.strftime("%d.%m.%Y"), t.created_at.strftime("%H:%M"),
            t.project.name, t.requester_name, t.requester_phone, t.channel_label, t.source_label,
            t.category.name, t.priority_label, t.description, t.created_by.full_name,
            t.duration_minutes, t.resolution, t.status_label,
        ])
        ws.cell(row=i, column=11).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=i, column=14).alignment = Alignment(wrap_text=True, vertical="top")
    for idx, (_, width) in enumerate(EXPORT_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(EXPORT_COLUMNS))}{max(len(tickets) + 1, 1)}"

    # Özet sayfası: aynı filtrelerle hesaplanan istatistikler
    ss = wb.create_sheet("Özet")
    s_conds = conds if filters.archived else stat_conditions(filters, user)
    s = summary(s_conds)
    rows = [
        ("Rapor tarihi", datetime.now().strftime("%d.%m.%Y %H:%M")),
        ("Filtreler", filter_description(filters)),
        (None, None),
        ("Toplam destek kaydı", s["total"]),
        ("Sistem Kaynaklı", s["system"]),
        ("Kullanıcı İşlemi Kaynaklı", s["user"]),
        ("Çözülen (Çözüldü + Kapatıldı)", s["done"]),
        ("Açık / Bekleyen", s["pending"]),
        ("Ortalama işlem süresi (dk)", s["avg_duration"]),
        ("Toplam işlem süresi (dk)", s["total_duration"]),
    ]
    for i, (k, v) in enumerate(rows, start=1):
        _write_row(ss, i, [k, v])
        if k:
            ss.cell(row=i, column=1).font = Font(bold=True)

    r = len(rows) + 2
    _header(ss, r, ["Proje", "Kayıt", "%"])
    for p in by_project(s_conds):
        r += 1
        _write_row(ss, r, [p["name"], p["count"], p["pct"]])
    r += 2
    _header(ss, r, ["Sorun Kategorisi", "Kayıt", "%"])
    for c in by_category(s_conds):
        r += 1
        _write_row(ss, r, [c["name"], c["count"], c["pct"]])
    r += 2
    _header(ss, r, ["Personel", "Kayıt", "Çözülen", "Bekleyen", "Toplam Süre (dk)"])
    for u in by_user(s_conds):
        r += 1
        _write_row(ss, r, [u["name"], u["total"], u["done"], u["pending"], u["duration"]])
    ss.column_dimensions["A"].width = 34
    for col in "BCDE":
        ss.column_dimensions[col].width = 16
    ss.cell(row=2, column=2).alignment = Alignment(wrap_text=True)
    return wb, len(tickets)


@bp.route("/export")
@admin_required
def export():
    filters = TicketFilters.from_args(request.args)
    wb, _ = build_workbook(filters, current_user)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"destek_kayitlari_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(
        buf, as_attachment=True, download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
