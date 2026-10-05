"""Gereksinimlerdeki 12 test senaryosu ve ek iş kuralı testleri.

Çalıştırma:  python -m pytest -q
"""
import re
from datetime import date, datetime, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app import create_app
from app import auth as auth_module
from app.constants import DONE_STATUSES, PENDING_STATUSES, ROLE_ADMIN, ROLE_STAFF
from app.demo import clear_demo, seed_demo
from app.extensions import db
from app.models import AuditLog, Category, Project, SupportTicket as T, User
from app.queries import TicketFilters, by_project, by_user, stat_conditions, summary
from app.tickets import create_ticket

PASSWORD = "Test.12345"


# --------------------------------------------------------------------------- #
# Kurulum
# --------------------------------------------------------------------------- #

@pytest.fixture
def app(tmp_path):
    app = create_app({
        "TESTING": True,
        "SECRET_KEY": "test",
        "SQLALCHEMY_DATABASE_URI": f"sqlite:///{tmp_path / 'test.db'}",
    }, data_dir=str(tmp_path))
    auth_module._failed_logins.clear()
    with app.app_context():
        for username, name, role in (
            ("yonetici", "Yönetici Kullanıcı", ROLE_ADMIN),
            ("personel1", "Personel Bir", ROLE_STAFF),
            ("personel2", "Personel İki", ROLE_STAFF),
        ):
            u = User(username=username, full_name=name, role=role)
            u.set_password(PASSWORD)
            db.session.add(u)
        db.session.commit()
    yield app


@pytest.fixture
def client(app):
    return app.test_client()


def csrf(client):
    """Oturumdaki CSRF jetonunu döndürür (yoksa bir sayfa açarak oluşturur)."""
    with client.session_transaction() as s:
        token = s.get("csrf_token")
    if not token:
        client.get("/login", follow_redirects=True)   # oturum açıksa ana sayfaya yönlenir
        with client.session_transaction() as s:
            token = s["csrf_token"]
    return token


def login(client, username, password=PASSWORD):
    client.get("/login")
    return client.post("/login", data={"csrf_token": csrf(client), "username": username, "password": password})


def ids(app):
    with app.app_context():
        return {
            "admin": db.session.execute(db.select(User.id).filter_by(username="yonetici")).scalar(),
            "p1": db.session.execute(db.select(User.id).filter_by(username="personel1")).scalar(),
            "p2": db.session.execute(db.select(User.id).filter_by(username="personel2")).scalar(),
            "projects": db.session.execute(db.select(Project.id).order_by(Project.id)).scalars().all(),
            "categories": db.session.execute(db.select(Category.id).order_by(Category.id)).scalars().all(),
        }


def ticket_form(app, **overrides):
    i = ids(app)
    data = {
        "project_id": i["projects"][0],
        "requester_name": "Ahmet Yılmaz",
        "requester_phone": "0312 555 44 33",
        "channel": "phone",
        "source": "system",
        "category_id": i["categories"][0],
        "description": "Kullanıcı doğru bilgilerle giriş yapmasına rağmen sisteme erişemiyor.",
        "priority": "normal",
        "duration_minutes": "8",
        "resolution": "Kontroller yapıldı ve problem giderildi.",
        "status": "resolved",
    }
    data.update({k: str(v) for k, v in overrides.items()})
    return data


def post_ticket(client, app, **overrides):
    data = ticket_form(app, **overrides)
    data["csrf_token"] = csrf(client)
    return client.post("/tickets/new", data=data)


def make_ticket(app, username, created_at=None, **overrides):
    """Doğrudan servis katmanı ile kayıt oluşturur (istatistik testleri için)."""
    with app.app_context():
        user = db.session.execute(db.select(User).filter_by(username=username)).scalar_one()
        data = ticket_form(app, **overrides)
        data = {
            "project_id": int(data["project_id"]), "category_id": int(data["category_id"]),
            "requester_name": data["requester_name"], "requester_phone": data["requester_phone"],
            "channel": data["channel"], "source": data["source"], "priority": data["priority"],
            "status": data["status"], "description": data["description"], "resolution": data["resolution"],
            "duration_minutes": int(data["duration_minutes"]) if data["duration_minutes"] not in ("", "None") else None,
        }
        return create_ticket(data, user, created_at=created_at).id


def page_text(response):
    return response.get_data(as_text=True)


# --------------------------------------------------------------------------- #
# Test 1: Personel giriş yapabiliyor mu?
# --------------------------------------------------------------------------- #

def test_01_staff_can_login(client):
    r = login(client, "personel1")
    assert r.status_code == 302
    r = client.get("/")
    assert r.status_code == 200
    assert "Personel Bir" in page_text(r)


def test_01b_wrong_password_and_inactive_user_rejected(client, app):
    assert login(client, "personel1", "yanlis-sifre").status_code == 401
    with app.app_context():
        u = db.session.execute(db.select(User).filter_by(username="personel2")).scalar_one()
        u.active = False
        db.session.commit()
    assert login(client, "personel2").status_code == 401
    # Oturum açmadan sayfalara erişilemez
    assert client.get("/tickets/").status_code == 302


def test_01c_passwords_are_hashed(app):
    with app.app_context():
        u = db.session.execute(db.select(User).filter_by(username="personel1")).scalar_one()
        assert u.password_hash != PASSWORD and PASSWORD not in u.password_hash


def test_01d_login_lockout_after_failed_attempts(client):
    for _ in range(5):
        login(client, "personel1", "hatali")
    assert login(client, "personel1").status_code == 429


def test_01e_post_without_csrf_token_is_rejected(client):
    r = client.post("/login", data={"username": "personel1", "password": PASSWORD})
    assert r.status_code == 400


# --------------------------------------------------------------------------- #
# Test 2-4: Kayıt oluşturma (sistem / kullanıcı işlemi kaynaklı)
# --------------------------------------------------------------------------- #

def test_02_create_ticket(client, app):
    login(client, "personel1")
    r = post_ticket(client, app)
    assert r.status_code == 302
    with app.app_context():
        t = db.session.execute(db.select(T)).scalar_one()
        assert re.fullmatch(rf"{date.today().year}-00001", t.ticket_no)
        assert t.created_by.username == "personel1"
        assert t.created_at.date() == date.today()
        assert t.duration_minutes == 8 and t.resolved_at is not None
        log = db.session.execute(db.select(AuditLog).filter_by(entity_type="ticket", entity_id=t.id)).scalar_one()
        assert log.action == "create" and log.user.username == "personel1"
    # Numaralar sıralı artar
    post_ticket(client, app)
    with app.app_context():
        assert db.session.execute(db.select(T.ticket_no).order_by(T.id.desc())).scalars().first().endswith("-00002")


def test_02b_validation_errors(client, app):
    login(client, "personel1")
    r = post_ticket(client, app, requester_name="", source="", requester_phone="abc")
    assert r.status_code == 400
    text = page_text(r)
    assert "Talep sahibinin adını giriniz" in text and "Sorun kaynağı seçiniz" in text
    # Çözüldü durumunda işlem süresi zorunlu
    r = post_ticket(client, app, duration_minutes="", status="resolved")
    assert r.status_code == 400 and "işlem süresini giriniz" in page_text(r)
    # Açık kayıtta süre girilmeyebilir
    assert post_ticket(client, app, duration_minutes="", status="open").status_code == 302


def test_03_create_system_sourced_ticket(client, app):
    login(client, "personel1")
    post_ticket(client, app, source="system")
    with app.app_context():
        assert db.session.execute(db.select(T)).scalar_one().source_label == "Sistem Kaynaklı"


def test_04_create_user_sourced_ticket(client, app):
    login(client, "personel1")
    post_ticket(client, app, source="user")
    with app.app_context():
        assert db.session.execute(db.select(T)).scalar_one().source_label == "Kullanıcı İşlemi Kaynaklı"


# --------------------------------------------------------------------------- #
# Test 5: Filtreleme
# --------------------------------------------------------------------------- #

def test_05_filter_tickets(client, app):
    i = ids(app)
    yesterday = datetime.now() - timedelta(days=1)
    make_ticket(app, "personel1", source="system", requester_name="Şükrü Işık")
    make_ticket(app, "personel1", source="user", project_id=i["projects"][1], status="open", duration_minutes="")
    make_ticket(app, "personel2", source="user", created_at=yesterday)
    login(client, "yonetici")

    def count(**params):
        text = page_text(client.get("/tickets/", query_string=params))
        return int(re.search(r'<p class="muted">(\d+) kayıt', text).group(1))

    assert count() == 3
    assert count(source="system") == 1
    assert count(source="user") == 2
    assert count(project_id=i["projects"][1]) == 1
    assert count(user_id=i["p2"]) == 1
    assert count(status="pending") == 1
    assert count(status="resolved") == 2
    assert count(date_from=date.today().isoformat(), date_to=date.today().isoformat()) == 2
    assert count(q="şükrü") == 1          # Türkçe büyük/küçük harf duyarsız arama
    assert count(q="IŞIK") == 1
    assert count(q="olmayan-ifade") == 0


# --------------------------------------------------------------------------- #
# Test 6-9: İstatistiklerin doğruluğu (demo veri üzerinde)
# --------------------------------------------------------------------------- #

def _all_tickets(app):
    with app.app_context():
        return [
            (t.source, t.status, t.project_id, t.created_by_id, t.duration_minutes, t.created_at.date())
            for t in db.session.execute(db.select(T).where(T.archived.is_(False))).scalars()
        ]


def test_06_dashboard_stats_match_records(client, app):
    with app.app_context():
        seed_demo(count=45)
    rows = _all_tickets(app)
    today_rows = [r for r in rows if r[5] == date.today()]
    expected = {
        "total": len(today_rows),
        "done": sum(r[1] in DONE_STATUSES for r in today_rows),
        "pending": sum(r[1] in PENDING_STATUSES for r in today_rows),
        "system": sum(r[0] == "system" for r in today_rows),
        "user": sum(r[0] == "user" for r in today_rows),
    }
    assert expected["total"] >= 6

    with app.app_context():
        admin = db.session.execute(db.select(User).filter_by(username="yonetici")).scalar_one()
        s = summary(stat_conditions(TicketFilters(date_from=date.today(), date_to=date.today()), admin))
        for key, value in expected.items():
            assert s[key] == value, key
        assert s["system"] + s["user"] == s["total"]

    login(client, "yonetici")
    html = page_text(client.get("/"))
    kpis = re.findall(r'<div class="kpi-value">([^<]+)</div>', html)
    assert [int(k) for k in kpis[:5]] == [expected[k] for k in ("total", "done", "pending", "system", "user")]


def test_07_project_stats(app):
    with app.app_context():
        seed_demo(count=45)
    rows = _all_tickets(app)
    with app.app_context():
        admin = db.session.execute(db.select(User).filter_by(username="yonetici")).scalar_one()
        result = {p["id"]: p["count"] for p in by_project(stat_conditions(TicketFilters(), admin))}
    for pid, count in result.items():
        assert count == sum(r[2] == pid for r in rows)
    assert sum(result.values()) == len(rows) == 45


def test_08_staff_stats(app):
    with app.app_context():
        seed_demo(count=45)
    rows = _all_tickets(app)
    with app.app_context():
        admin = db.session.execute(db.select(User).filter_by(username="yonetici")).scalar_one()
        result = by_user(stat_conditions(TicketFilters(), admin))
    assert sum(u["total"] for u in result) == 45
    for u in result:
        mine = [r for r in rows if r[3] == u["id"]]
        assert u["total"] == len(mine)
        assert u["done"] == sum(r[1] in DONE_STATUSES for r in mine)
        assert u["pending"] == sum(r[1] in PENDING_STATUSES for r in mine)
        assert u["duration"] == sum(r[4] or 0 for r in mine)


def test_09_average_duration(app):
    for minutes in (3, 7, 15, 25):
        make_ticket(app, "personel1", duration_minutes=minutes)
    make_ticket(app, "personel1", duration_minutes="", status="open")   # süresiz kayıt ortalamaya girmez
    with app.app_context():
        admin = db.session.execute(db.select(User).filter_by(username="yonetici")).scalar_one()
        s = summary(stat_conditions(TicketFilters(), admin))
    assert s["avg_duration"] == 12.5
    assert s["total_duration"] == 50
    assert s["total"] == 5


# --------------------------------------------------------------------------- #
# Test 10: Excel çıktısı
# --------------------------------------------------------------------------- #

def test_10_excel_export(client, app):
    make_ticket(app, "personel1", source="system", description="=HYPERLINK(\"http://x\")")
    make_ticket(app, "personel1", source="user")
    make_ticket(app, "personel2", source="user")
    login(client, "yonetici")
    r = client.get("/reports/export", query_string={"source": "user"})
    assert r.status_code == 200
    assert r.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    wb = load_workbook(BytesIO(r.data))
    ws = wb["Kayıtlar"]
    headers = [c.value for c in ws[1]]
    for col in ("Kayıt No", "Tarih", "Saat", "Proje", "Talep Sahibi", "Telefon", "Sorun Kaynağı",
                "Sorun Kategorisi", "Açıklama", "Personel", "İşlem Süresi (dk)", "Sonuç / Yapılan İşlem", "Durum"):
        assert col in headers, col
    assert ws.max_row == 3   # başlık + 2 filtrelenmiş kayıt
    assert {ws.cell(row=i, column=headers.index("Sorun Kaynağı") + 1).value for i in (2, 3)} == {"Kullanıcı İşlemi Kaynaklı"}
    # Grafikli rapor: ilk sayfa "Rapor" (özet kutuları + 5 grafik), sayılar Kayıtlar'dan formülle hesaplanır
    assert wb.sheetnames == ["Rapor", "Tablolar", "Kayıtlar"]
    assert wb.active.title == "Rapor"
    assert len(wb["Rapor"]._charts) == 5
    tables = wb["Tablolar"]
    assert tables["B2"].value == "=COUNTA('Kayıtlar'!$A$2:$A$3)"
    assert tables["B3"].value.startswith("=COUNTIF('Kayıtlar'!$H$2:$H$3,\"Sistem Kaynaklı\")")
    assert wb["Rapor"]["A7"].value == "='Tablolar'!$B$2"
    assert wb.calculation.fullCalcOnLoad   # Excel açılışta hesaplar

    # Formül enjeksiyonu: "=" ile başlayan metin formül değil, düz metin olarak yazılır
    wb_all = load_workbook(BytesIO(client.get("/reports/export").data))
    cells = [c for row in wb_all["Kayıtlar"].iter_rows(min_row=2) for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert cells and all(c.data_type == "s" for c in cells)


def test_10b_excel_report_empty_and_long_range(client, app):
    """Kayıt yokken ve uzun tarih aralığında (aylık grafik) rapor sorunsuz oluşur."""
    login(client, "yonetici")
    r = client.get("/reports/export", query_string={"date_from": "2026-01-01", "date_to": "2026-01-31"})
    wb = load_workbook(BytesIO(r.data))
    assert wb["Kayıtlar"].max_row == 1 and len(wb["Rapor"]._charts) == 5
    make_ticket(app, "personel1", created_at=datetime.now() - timedelta(days=70))
    make_ticket(app, "personel1")
    r = client.get("/reports/export", query_string={"date_from": (date.today() - timedelta(days=90)).isoformat(),
                                                  "date_to": date.today().isoformat()})
    tables = load_workbook(BytesIO(r.data))["Tablolar"]
    assert any(row[0].value == "Aylara Göre" for row in tables.iter_rows(max_col=1))


def test_10c_monthly_report(client, app):
    """Aylık bilgilendirme raporu: ay ay Bireysel/Sistemsel tablo, dairesel grafikler, kişisel veri yok."""
    i = ids(app)
    year = date.today().year
    may, june = datetime(year, 5, 10, 10, 0), datetime(year, 6, 3, 9, 0)
    make_ticket(app, "personel1", created_at=may, source="user", requester_name="Gizli Kişi",
                requester_phone="0555 111 22 33")
    make_ticket(app, "personel1", created_at=may, source="system")
    make_ticket(app, "personel2", created_at=may, source="system", project_id=i["projects"][1])
    make_ticket(app, "personel2", created_at=june, source="user")
    archived = make_ticket(app, "personel2", created_at=june, source="system")
    with app.app_context():
        db.session.get(T, archived).archived = True
        db.session.commit()

    login(client, "yonetici")
    assert "Aylık Bilgilendirme Raporu" in page_text(client.get("/reports/"))
    r = client.get("/reports/monthly", query_string={"year": year, "month": 6})
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.data))
    ws = wb.active
    assert ws.title == f"Haziran {year}"
    cells = {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}
    text = " ".join(str(v) for v in cells.values())
    # KVKK: birimlere gönderilen dosyada kişisel veri bulunmaz
    assert "Gizli Kişi" not in text and "0555" not in text
    rows = {ws.cell(row=r, column=1).value: r for r in range(1, ws.max_row + 1)}
    # Sayfa 1 tablosu: Proje (A) | Bireysel (E) | Sistemsel (G) | Toplam (I) — arşivlenen dahil değil
    assert ws.cell(row=rows["Haziran Toplam"], column=5).value.startswith("=SUM(")
    project_row = rows["Dijital Otomasyon Projesi 1"]
    assert [ws.cell(row=project_row, column=c).value for c in (5, 7)] == [1, 0]
    # Sayfa 2 matrisi: Mayıs satırında 2. projenin Bireysel/Sistemsel değerleri (sütun E-F)
    assert [ws.cell(row=rows["Mayıs"], column=c).value for c in (5, 6)] == [0, 1]
    assert ws.cell(row=rows["Genel Toplam"], column=4).value.startswith("=SUM(")
    assert len(ws._charts) == 4      # 3 halka grafik + aylık eğilim grafiği
    # Geçersiz parametre
    assert client.get("/reports/monthly", query_string={"year": year, "month": 13}).status_code == 400


# --------------------------------------------------------------------------- #
# Test 11: Personel yönetici ekranlarına erişemiyor mu?
# --------------------------------------------------------------------------- #

def test_11_staff_cannot_access_admin_pages(client, app):
    other = make_ticket(app, "personel2")
    own = make_ticket(app, "personel1")
    login(client, "personel1")
    for url in ("/admin/", "/admin/projects", "/admin/categories", "/admin/users", "/admin/settings",
                "/reports/", "/reports/export", "/reports/monthly"):
        assert client.get(url).status_code == 403, url
    token = csrf(client)
    assert client.post("/admin/projects", data={"csrf_token": token, "name": "X"}).status_code == 403
    assert client.post(f"/tickets/{own}/archive", data={"csrf_token": token}).status_code == 403
    # Başka personelin kaydı görülemez / düzenlenemez
    assert client.get(f"/tickets/{other}").status_code == 404
    assert client.get(f"/tickets/{other}/edit").status_code == 404
    assert client.post(f"/tickets/{other}/status", data={"csrf_token": token, "status": "closed"}).status_code == 404
    # Listede yalnızca kendi kaydı var; Yönetim/Raporlar menüsü görünmez
    html = page_text(client.get("/tickets/"))
    assert "1 kayıt" in html
    assert "/admin/" not in html and "/reports/" not in html
    # Personel istatistiklerinde personel iş yükü tablosu yok
    assert "Personel İstatistikleri" not in page_text(client.get("/stats/"))
    assert "Personel İş Yükü" not in page_text(client.get("/"))


# --------------------------------------------------------------------------- #
# Test 12: Yönetici tüm kayıtları görebiliyor mu?
# --------------------------------------------------------------------------- #

def test_12_admin_sees_all_tickets(client, app):
    t1 = make_ticket(app, "personel1")
    t2 = make_ticket(app, "personel2")
    login(client, "yonetici")
    html = page_text(client.get("/tickets/"))
    assert "2 kayıt" in html
    assert client.get(f"/tickets/{t1}").status_code == 200
    assert client.get(f"/tickets/{t2}").status_code == 200
    assert client.get("/admin/projects").status_code == 200


# --------------------------------------------------------------------------- #
# Ek iş kuralları
# --------------------------------------------------------------------------- #

def test_status_change_reopen_and_audit(client, app):
    tid = make_ticket(app, "personel1", status="open", duration_minutes="")
    login(client, "personel1")
    token = csrf(client)
    # Süre girilmeden çözüldü yapılamaz
    client.post(f"/tickets/{tid}/status", data={"csrf_token": token, "status": "resolved"})
    with app.app_context():
        assert db.session.get(T, tid).status == "open"
    client.post(f"/tickets/{tid}/status", data={"csrf_token": token, "status": "resolved", "duration_minutes": "6"})
    with app.app_context():
        t = db.session.get(T, tid)
        assert t.status == "resolved" and t.duration_minutes == 6 and t.resolved_at
    # Yeniden açma
    client.post(f"/tickets/{tid}/status", data={"csrf_token": token, "status": "open"})
    with app.app_context():
        t = db.session.get(T, tid)
        assert t.status == "open" and t.resolved_at is None and t.updated_by.username == "personel1"
        actions = [a.action for a in db.session.execute(
            db.select(AuditLog).filter_by(entity_type="ticket", entity_id=tid).order_by(AuditLog.id)).scalars()]
        assert actions == ["create", "status", "status"]


def test_save_and_similar_keeps_fields_except_requester(client, app):
    """Tek personel hızlı giriş: kaydet ve benzerini gir -> sadece talep sahibi değişir."""
    i = ids(app)
    login(client, "personel1")
    r = post_ticket(client, app, source="user", category_id=i["categories"][2], project_id=i["projects"][1],
                    description="Toplu bildirim ekranında hata", duration_minutes=4, save_and_similar=1)
    assert r.status_code == 302 and "kopya=" in r.headers["Location"]
    html = page_text(client.get(r.headers["Location"]))
    assert "numaralı kayıttaki bilgiler aktarıldı" in html
    assert "Toplu bildirim ekranında hata" in html                          # açıklama kopyalandı
    assert 'name="requester_name" value=""' in html                         # talep sahibi boş
    assert 'name="requester_phone" value=""' in html
    assert f'name="project_id" value="{i["projects"][1]}" checked' in html  # proje seçili
    # Yalnızca adı değiştirip ikinci kayıt
    data = ticket_form(app, source="user", category_id=i["categories"][2], project_id=i["projects"][1],
                       description="Toplu bildirim ekranında hata", duration_minutes=4,
                       requester_name="İkinci Kişi", requester_phone="")
    data["csrf_token"] = csrf(client)
    assert client.post("/tickets/new", data=data).status_code == 302
    with app.app_context():
        a, b = db.session.execute(db.select(T).order_by(T.id)).scalars().all()
        assert (a.project_id, a.category_id, a.source, a.description) == (b.project_id, b.category_id, b.source, b.description)
        assert b.requester_name == "İkinci Kişi" and b.requester_phone is None and b.ticket_no != a.ticket_no


def test_copy_from_detail_respects_ownership(client, app):
    own = make_ticket(app, "personel1")
    other = make_ticket(app, "personel2")
    login(client, "personel1")
    assert "Benzer Kayıt Oluştur" in page_text(client.get(f"/tickets/{own}"))
    assert client.get("/tickets/new", query_string={"kopya": own}).status_code == 200
    assert client.get("/tickets/new", query_string={"kopya": other}).status_code == 404   # başkasının kaydı


def test_edit_ticket_records_changes(client, app):
    tid = make_ticket(app, "personel1")
    login(client, "personel1")
    data = ticket_form(app, source="user", requester_name="Yeni İsim")
    data["csrf_token"] = csrf(client)
    assert client.post(f"/tickets/{tid}/edit", data=data).status_code == 302
    with app.app_context():
        t = db.session.get(T, tid)
        assert t.source == "user" and t.requester_name == "Yeni İsim"
        log = db.session.execute(db.select(AuditLog).filter_by(entity_id=tid, action="update")).scalar_one()
        assert "Sorun kaynağı" in log.changes_dict and "Talep sahibi" in log.changes_dict


def test_archive_excludes_from_stats_and_no_delete(client, app):
    tid = make_ticket(app, "personel1")
    make_ticket(app, "personel1")
    login(client, "yonetici")
    client.post(f"/tickets/{tid}/archive", data={"csrf_token": csrf(client)})
    with app.app_context():
        admin = db.session.execute(db.select(User).filter_by(username="yonetici")).scalar_one()
        assert db.session.get(T, tid).archived is True       # silinmedi, arşivlendi
        assert summary(stat_conditions(TicketFilters(), admin))["total"] == 1
    assert "1 kayıt" in page_text(client.get("/tickets/"))
    assert "1 kayıt" in page_text(client.get("/tickets/?archived=1"))
    client.post(f"/tickets/{tid}/archive", data={"csrf_token": csrf(client), "restore": "1"})
    assert "2 kayıt" in page_text(client.get("/tickets/"))


def test_admin_manages_projects_categories_users(client, app):
    login(client, "yonetici")
    token = csrf(client)
    client.post("/admin/projects", data={"csrf_token": token, "name": "Dijital Otomasyon Projesi 4"})
    with app.app_context():
        p = db.session.execute(db.select(Project).filter_by(name="Dijital Otomasyon Projesi 4")).scalar_one()
        pid = p.id
    client.post(f"/admin/projects/{pid}", data={"csrf_token": token, "name": "Yeni Ad", "sort_order": "9", "active": "0"})
    with app.app_context():
        p = db.session.get(Project, pid)
        assert p.name == "Yeni Ad" and not p.active
    # Pasif proje yeni kayıt formunda seçilemez
    login(client, "personel1")
    r = post_ticket(client, app, project_id=pid)
    assert r.status_code == 400

    login(client, "yonetici")
    token = csrf(client)
    client.post("/admin/categories", data={"csrf_token": token, "name": "Raporlama"})
    client.post("/admin/users", data={"csrf_token": token, "username": "yeni.personel", "full_name": "Yeni Personel",
                                      "role": "staff", "password": "Sifre.123", "password2": "Sifre.123"})
    with app.app_context():
        assert db.session.execute(db.select(Category).filter_by(name="Raporlama")).scalar_one()
        assert db.session.execute(db.select(User).filter_by(username="yeni.personel")).scalar_one().role == "staff"
    assert login(client, "yeni.personel", "Sifre.123").status_code == 302


def test_admin_cannot_demote_self(client, app):
    login(client, "yonetici")
    admin_id = ids(app)["admin"]
    client.post(f"/admin/users/{admin_id}", data={"csrf_token": csrf(client), "full_name": "X", "role": "staff", "active": "1"})
    with app.app_context():
        assert db.session.get(User, admin_id).role == ROLE_ADMIN


def test_demo_data_can_be_cleared(app):
    real = make_ticket(app, "personel1")
    with app.app_context():
        seed_demo(count=30)
        assert db.session.execute(db.select(db.func.count(T.id))).scalar() == 31
        tickets, users, _ = clear_demo()
        assert tickets == 30 and users == 5
        assert db.session.execute(db.select(T.id)).scalars().all() == [real]


def test_all_pages_render(client, app):
    with app.app_context():
        seed_demo(count=20)
    login(client, "yonetici")
    for url in ("/", "/?period=week", "/?period=month", "/tickets/", "/tickets/new", "/tickets/1", "/tickets/1/edit",
                "/stats/", "/stats/?period=day", "/stats/?period=week", "/stats/?period=custom&date_from=2020-01-01&date_to=2030-01-01",
                "/reports/", "/admin/projects", "/admin/categories", "/admin/users", "/admin/settings", "/password"):
        assert client.get(url).status_code == 200, url


# --------------------------------------------------------------------------- #
# Tek bilgisayar kurulumu: ilk kurulum ekranı ve yedekleme
# --------------------------------------------------------------------------- #

@pytest.fixture
def fresh_client(tmp_path):
    """Hiç kullanıcısı olmayan, yeni kurulmuş program."""
    app = create_app({"TESTING": True, "SECRET_KEY": "test",
                      "SQLALCHEMY_DATABASE_URI": f"sqlite:///{tmp_path / 'fresh.db'}"}, data_dir=str(tmp_path))
    return app, app.test_client()


def test_initial_setup_creates_admin(fresh_client):
    app, client = fresh_client
    # Yönetici yokken her sayfa kurulum ekranına yönlenir
    r = client.get("/tickets/")
    assert r.status_code == 302 and r.headers["Location"].endswith("/kurulum")
    assert "İlk Kurulum" in page_text(client.get("/kurulum"))
    r = client.post("/kurulum", data={"csrf_token": csrf(client), "org_name": "Bilgi İşlem",
                                      "full_name": "Ayşe Yönetici", "username": "admin",
                                      "password": "Kurulum.123", "password2": "Kurulum.123", "demo": "1"})
    assert r.status_code == 302
    assert client.get("/admin/settings").status_code == 200     # otomatik giriş yapıldı
    with app.app_context():
        admin = db.session.execute(db.select(User).filter_by(username="admin")).scalar_one()
        assert admin.role == ROLE_ADMIN and admin.check_password("Kurulum.123")
        assert db.session.execute(db.select(db.func.count(T.id))).scalar() == 45
    # Kurulum tamamlandıktan sonra kurulum ekranı tekrar kullanılamaz
    assert client.get("/kurulum").status_code == 302


def test_remote_computers_are_rejected(fresh_client, client):
    """Program yalnızca kurulu olduğu bilgisayardan kullanılır; ağdan gelen istekler reddedilir."""
    app, fresh = fresh_client
    assert fresh.get("/kurulum", environ_base={"REMOTE_ADDR": "192.168.1.50"}).status_code == 403
    assert fresh.post("/kurulum", environ_base={"REMOTE_ADDR": "192.168.1.50"},
                      data={"full_name": "X", "username": "x", "password": "Kurulum.123",
                            "password2": "Kurulum.123"}).status_code == 403
    with app.app_context():
        assert db.session.execute(db.select(db.func.count(User.id))).scalar() == 0
    # Kurulu programda da giriş dahil hiçbir sayfa ağdan açılamaz
    for url in ("/login", "/", "/tickets/", "/static/css/app.css"):
        assert client.get(url, environ_base={"REMOTE_ADDR": "10.0.0.7"}).status_code == 403, url
    assert client.get("/login", environ_base={"REMOTE_ADDR": "::1"}).status_code == 200


def test_idle_session_is_logged_out(client, app):
    login(client, "personel1")
    assert client.get("/").status_code == 200
    assert client.get("/oturum").status_code == 204          # aktif kullanımda oturum açık kalır
    with client.session_transaction() as s:
        s["last_seen"] -= app.config["IDLE_TIMEOUT_MINUTES"] * 60 + 1
    r = client.get("/tickets/")
    assert r.status_code == 302 and "/login" in r.headers["Location"]
    assert client.get("/tickets/").status_code == 302          # oturum gerçekten kapandı


def test_backup(app, client):
    from app.backup import create_backup, list_backups
    make_ticket(app, "personel1")
    path = create_backup(app, daily=True)
    assert path and create_backup(app, daily=True) is None      # günde bir otomatik yedek
    import sqlite3
    with sqlite3.connect(path) as con:
        assert con.execute("select count(*) from support_tickets").fetchone()[0] == 1
    login(client, "yonetici")
    client.post("/admin/settings", data={"csrf_token": csrf(client), "action": "backup_now"})
    assert len(list_backups(app)) == 2
